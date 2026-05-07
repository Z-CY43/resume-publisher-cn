#!/usr/bin/env python
"""Generate synchronized A4 HTML, DOCX, and PDF resumes from YAML + Markdown sources."""

from __future__ import annotations

import argparse
import base64
import html
import json
import os
import re
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


DEFAULT_ACCENT = "#1F4E79"
SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_CSS = SKILL_DIR / "assets" / "templates" / "a4-resume.css"


class ResumeSourceError(ValueError):
    pass


def strip_comment(line: str) -> str:
    quote: str | None = None
    escaped = False
    for idx, char in enumerate(line):
        if char == "\\" and quote and not escaped:
            escaped = True
            continue
        if char in ("'", '"') and not escaped:
            quote = None if quote == char else char if quote is None else quote
        if char == "#" and quote is None:
            return line[:idx].rstrip()
        escaped = False
    return line.rstrip()


def parse_scalar(raw: str) -> Any:
    raw = raw.strip()
    if raw == "":
        return ""
    if raw in ("null", "Null", "NULL", "~"):
        return None
    if raw in ("true", "True", "TRUE"):
        return True
    if raw in ("false", "False", "FALSE"):
        return False
    if raw.startswith("[") and raw.endswith("]"):
        inner = raw[1:-1].strip()
        if not inner:
            return []
        parts: list[str] = []
        current: list[str] = []
        quote: str | None = None
        for char in inner:
            if char in ("'", '"'):
                quote = None if quote == char else char if quote is None else quote
            if char == "," and quote is None:
                parts.append("".join(current).strip())
                current = []
            else:
                current.append(char)
        parts.append("".join(current).strip())
        return [parse_scalar(part) for part in parts]
    if (raw.startswith('"') and raw.endswith('"')) or (raw.startswith("'") and raw.endswith("'")):
        try:
            return json.loads(raw) if raw.startswith('"') else raw[1:-1]
        except json.JSONDecodeError:
            return raw[1:-1]
    if re.fullmatch(r"-?\d+", raw):
        try:
            return int(raw)
        except ValueError:
            pass
    if re.fullmatch(r"-?\d+\.\d+", raw):
        try:
            return float(raw)
        except ValueError:
            pass
    return raw


def load_yaml_like(path: Path | None) -> dict[str, Any]:
    if path is None or not path.exists():
        return {}
    text = path.read_text(encoding="utf-8-sig")
    try:
        import yaml  # type: ignore

        loaded = yaml.safe_load(text)
        return loaded or {}
    except Exception:
        pass

    raw_lines = text.splitlines()
    lines: list[tuple[int, str]] = []
    for line in raw_lines:
        clean = strip_comment(line)
        if not clean.strip():
            continue
        indent = len(clean) - len(clean.lstrip(" "))
        lines.append((indent, clean.strip()))
    if not lines:
        return {}

    def parse_block(index: int, indent: int) -> tuple[Any, int]:
        if index >= len(lines) or lines[index][0] < indent:
            return {}, index
        if lines[index][0] == indent and lines[index][1].startswith("- "):
            return parse_list(index, indent)
        return parse_dict(index, indent)

    def collect_multiline(index: int, parent_indent: int, folded: bool) -> tuple[str, int]:
        collected: list[str] = []
        while index < len(lines) and lines[index][0] > parent_indent:
            collected.append(lines[index][1])
            index += 1
        if folded:
            return " ".join(collected), index
        return "\n".join(collected), index

    def parse_key_value(text_line: str, index: int, indent: int) -> tuple[str, Any, int]:
        if ":" not in text_line:
            return text_line, "", index + 1
        key, rest = text_line.split(":", 1)
        key = key.strip().strip("'\"")
        rest = rest.strip()
        if rest in ("|", "|-", ">", ">-"):
            value, next_index = collect_multiline(index + 1, indent, rest.startswith(">"))
            return key, value, next_index
        if rest == "":
            value, next_index = parse_block(index + 1, indent + 2)
            return key, value, next_index
        return key, parse_scalar(rest), index + 1

    def parse_dict(index: int, indent: int) -> tuple[dict[str, Any], int]:
        result: dict[str, Any] = {}
        while index < len(lines):
            current_indent, text_line = lines[index]
            if current_indent < indent:
                break
            if current_indent > indent:
                index += 1
                continue
            if text_line.startswith("- "):
                break
            key, value, index = parse_key_value(text_line, index, indent)
            result[key] = value
        return result, index

    def parse_list(index: int, indent: int) -> tuple[list[Any], int]:
        result: list[Any] = []
        while index < len(lines):
            current_indent, text_line = lines[index]
            if current_indent < indent:
                break
            if current_indent != indent or not text_line.startswith("- "):
                break
            rest = text_line[2:].strip()
            if rest == "":
                value, index = parse_block(index + 1, indent + 2)
                result.append(value)
                continue
            if ":" in rest and not rest.startswith(("http://", "https://")):
                key, value, next_index = parse_key_value(rest, index, indent)
                item: dict[str, Any] = {key: value}
                if next_index < len(lines) and lines[next_index][0] > indent:
                    nested, next_index = parse_dict(next_index, indent + 2)
                    if isinstance(nested, dict):
                        item.update(nested)
                result.append(item)
                index = next_index
                continue
            result.append(parse_scalar(rest))
            index += 1
        return result, index

    parsed, final_index = parse_block(0, lines[0][0])
    if final_index < len(lines):
        raise ResumeSourceError(f"Could not parse YAML near: {lines[final_index][1]}")
    if not isinstance(parsed, dict):
        raise ResumeSourceError("Top-level YAML source must be a mapping.")
    return parsed


def as_list(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, list):
        return value
    return [value]


def text_value(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


@dataclass
class Entry:
    heading: str
    bullets: list[str] = field(default_factory=list)
    paragraphs: list[str] = field(default_factory=list)


@dataclass
class Section:
    title: str
    items: list[str] = field(default_factory=list)
    paragraphs: list[str] = field(default_factory=list)
    entries: list[Entry] = field(default_factory=list)


@dataclass
class Resume:
    name: str
    target_role: str
    contact: list[str]
    accent: str
    photo: Path | None
    sections: list[Section]


def normalize_entry(raw: Any) -> Entry:
    if isinstance(raw, str):
        return Entry(heading=raw)
    if not isinstance(raw, dict):
        return Entry(heading=text_value(raw))
    heading_parts = [raw.get("heading"), raw.get("title"), raw.get("name"), raw.get("company"), raw.get("school")]
    heading = next((text_value(part) for part in heading_parts if text_value(part)), "")
    role = text_value(raw.get("role") or raw.get("position"))
    date = text_value(raw.get("date") or raw.get("period"))
    if role and role not in heading:
        heading = f"{heading} | {role}" if heading else role
    if date and date not in heading:
        heading = f"{heading}  {date}" if heading else date
    bullets = [text_value(item) for item in as_list(raw.get("bullets") or raw.get("items")) if text_value(item)]
    paragraphs = [text_value(item) for item in as_list(raw.get("paragraphs") or raw.get("description")) if text_value(item)]
    return Entry(heading=heading, bullets=bullets, paragraphs=paragraphs)


def normalize_section(raw: Any) -> Section | None:
    if isinstance(raw, str):
        return Section(title=raw)
    if not isinstance(raw, dict):
        return None
    title = text_value(raw.get("title") or raw.get("name"))
    if not title:
        return None
    section = Section(title=title)
    section.items = [text_value(item) for item in as_list(raw.get("items") or raw.get("bullets")) if text_value(item)]
    section.paragraphs = [text_value(item) for item in as_list(raw.get("paragraphs") or raw.get("description")) if text_value(item)]
    section.entries = [normalize_entry(item) for item in as_list(raw.get("entries"))]
    section.entries = [entry for entry in section.entries if entry.heading or entry.bullets or entry.paragraphs]
    return section


def parse_markdown_sections(path: Path | None) -> list[Section]:
    if path is None or not path.exists():
        return []
    sections: list[Section] = []
    current: Section | None = None
    for raw_line in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            current = Section(title=heading.group(2).strip())
            sections.append(current)
            continue
        if current is None:
            current = Section(title="补充信息")
            sections.append(current)
        if line.startswith(("- ", "* ")):
            current.items.append(line[2:].strip())
        else:
            current.paragraphs.append(line)
    return sections


def build_conventional_sections(data: dict[str, Any]) -> list[Section]:
    mapping = [
        ("个人优势", ["advantages", "highlights", "summary", "个人优势"]),
        ("核心技能", ["skills", "core_skills", "技能", "核心技能"]),
        ("工作经历", ["experiences", "work_experience", "工作经历"]),
        ("重点项目", ["projects", "project_experience", "项目经历", "重点项目"]),
        ("校园/个人项目", ["campus_projects", "personal_projects", "校园项目", "个人项目"]),
        ("竞赛/荣誉", ["awards", "honors", "competitions", "奖项", "竞赛经历"]),
        ("教育背景", ["education", "教育背景"]),
        ("自我评价", ["self_evaluation", "自我评价"]),
    ]
    sections: list[Section] = []
    for title, keys in mapping:
        value = next((data[key] for key in keys if key in data and data[key]), None)
        if value is None:
            continue
        if isinstance(value, list) and value and all(isinstance(item, dict) for item in value):
            entries = [normalize_entry(item) for item in value]
            sections.append(Section(title=title, entries=[entry for entry in entries if entry.heading or entry.bullets]))
        else:
            sections.append(Section(title=title, items=[text_value(item) for item in as_list(value) if text_value(item)]))
    return sections


def contact_items(data: dict[str, Any]) -> list[str]:
    contact = data.get("contact") or data.get("contacts") or data.get("联系方式")
    if isinstance(contact, dict):
        ordered_keys = ["phone", "email", "wechat", "location", "github", "portfolio"]
        items = [text_value(contact.get(key)) for key in ordered_keys if text_value(contact.get(key))]
        extras = [text_value(value) for key, value in contact.items() if key not in ordered_keys and text_value(value)]
        return items + extras
    return [text_value(item) for item in as_list(contact) if text_value(item)]


def resolve_photo(data: dict[str, Any]) -> Path | None:
    raw = text_value(
        data.get("photo")
        or data.get("avatar")
        or data.get("portrait")
        or data.get("证件照")
        or data.get("照片")
    )
    if not raw:
        return None
    photo = Path(raw)
    if not photo.is_absolute():
        source_dir = Path(text_value(data.get("_source_dir")) or ".")
        photo = source_dir / photo
    if not photo.exists():
        raise ResumeSourceError(f"Photo file not found: {photo}")
    return photo


def build_resume(data: dict[str, Any], markdown_sections: list[Section]) -> Resume:
    name = text_value(data.get("name") or data.get("姓名") or data.get("candidate"))
    target_role = text_value(data.get("target_role") or data.get("role") or data.get("position") or data.get("求职方向") or data.get("岗位"))
    theme = data.get("theme") if isinstance(data.get("theme"), dict) else {}
    accent = text_value(theme.get("accent") or data.get("accent") or data.get("theme_color")) or DEFAULT_ACCENT

    sections = [section for section in (normalize_section(item) for item in as_list(data.get("sections"))) if section]
    if not sections:
        sections = build_conventional_sections(data)
    sections.extend(markdown_sections)

    if not name:
        raise ResumeSourceError("resume.yaml must include name.")
    if not target_role:
        target_role = "求职简历"
    if not sections:
        raise ResumeSourceError("No resume sections found. Add sections to resume.yaml or resume.md.")

    return Resume(
        name=name,
        target_role=target_role,
        contact=contact_items(data),
        accent=accent,
        photo=resolve_photo(data),
        sections=sections,
    )


def safe_filename(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", value)
    value = re.sub(r"\s+", "", value)
    return value or "resume"


def default_stem(resume: Resume) -> str:
    return f"{safe_filename(resume.name)}-{safe_filename(resume.target_role)}-简历"


def image_data_uri(path: Path) -> str:
    suffix = path.suffix.lower()
    mime = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
        ".gif": "image/gif",
    }.get(suffix, "application/octet-stream")
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime};base64,{encoded}"


def render_html(resume: Resume, output: Path, css_path: Path = DEFAULT_CSS) -> None:
    css = css_path.read_text(encoding="utf-8") if css_path.exists() else ""
    css = css.replace("--accent: #1F4E79;", f"--accent: {resume.accent};")
    title = f"{resume.name}-{resume.target_role}-简历"
    parts = [
        "<!doctype html>",
        '<html lang="zh-CN">',
        "<head>",
        '  <meta charset="utf-8">',
        '  <meta name="viewport" content="width=device-width, initial-scale=1">',
        f"  <title>{html.escape(title)}</title>",
        "  <style>",
        css,
        "  </style>",
        "</head>",
        "<body>",
        '  <main class="page">',
        f'    <header class="resume-header{" has-photo" if resume.photo else ""}">',
        '      <div class="identity">',
        f"        <h1>{html.escape(resume.name)}</h1>",
        f"        <div class=\"target\">{html.escape(resume.target_role)}</div>",
        "      </div>",
    ]
    if resume.contact:
        parts.append('      <div class="contact">')
        for item in resume.contact:
            parts.append(f"        <span>{html.escape(item)}</span>")
        parts.append("      </div>")
    if resume.photo:
        parts.append(
            f'      <img class="profile-photo" src="{html.escape(image_data_uri(resume.photo), quote=True)}" alt="{html.escape(resume.name)}证件照">'
        )
    parts.extend(["    </header>"])
    for section in resume.sections:
        parts.append("    <section>")
        parts.append(f"      <h2>{html.escape(section.title)}</h2>")
        for paragraph in section.paragraphs:
            parts.append(f"      <p>{html.escape(paragraph)}</p>")
        if section.items:
            parts.append("      <ul>")
            for item in section.items:
                parts.append(f"        <li>{html.escape(item)}</li>")
            parts.append("      </ul>")
        for entry in section.entries:
            if entry.heading:
                parts.append(f"      <h3>{html.escape(entry.heading)}</h3>")
            for paragraph in entry.paragraphs:
                parts.append(f"      <p>{html.escape(paragraph)}</p>")
            if entry.bullets:
                parts.append("      <ul>")
                for bullet in entry.bullets:
                    parts.append(f"        <li>{html.escape(bullet)}</li>")
                parts.append("      </ul>")
        parts.append("    </section>")
    parts.extend(["  </main>", "</body>", "</html>", ""])
    output.write_text("\n".join(parts), encoding="utf-8")


def set_run_font(run: Any, font_name: str, size_pt: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if color:
        color = color.lstrip("#")
        if len(color) == 6:
            run.font.color.rgb = RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))


def style_paragraph(paragraph: Any, font_name: str = "Microsoft YaHei", size_pt: float = 10.5) -> None:
    from docx.shared import Pt

    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.08
    for run in paragraph.runs:
        set_run_font(run, font_name, size_pt)


def add_colored_heading(document: Any, text: str, level: int, accent: str) -> Any:
    paragraph = document.add_paragraph()
    paragraph.style = f"Heading {level}"
    run = paragraph.add_run(text)
    size = 14 if level == 2 else 11.5
    set_run_font(run, "Microsoft YaHei", size, True, accent)
    return paragraph


def add_floating_photo_with_word(docx_path: Path, photo_path: Path) -> None:
    if os.name != "nt":
        raise ResumeSourceError("Floating DOCX photo placement requires Microsoft Word on Windows.")
    safe_docx = str(docx_path).replace("'", "''")
    safe_photo = str(photo_path).replace("'", "''")
    ps_script = f"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {{
  $doc = $word.Documents.Open('{safe_docx}')
  $pageWidth = $doc.PageSetup.PageWidth
  $rightMargin = $doc.PageSetup.RightMargin
  $topMargin = $doc.PageSetup.TopMargin
  $width = 68
  $height = 79
  $left = $pageWidth - $rightMargin - $width
  $top = $topMargin - 2
  $shape = $doc.Shapes.AddPicture('{safe_photo}', $false, $true, $left, $top, $width, $height)
  $shape.Name = 'resume-id-photo'
  $shape.AlternativeText = '证件照'
  $shape.RelativeHorizontalPosition = 1
  $shape.RelativeVerticalPosition = 1
  $shape.Left = $left
  $shape.Top = $top
  $shape.Width = $width
  $shape.Height = $height
  $shape.WrapFormat.Type = 6
  $doc.Save()
  $doc.Close([ref] $false)
}} finally {{
  $word.Quit()
}}
"""
    encoded = base64.b64encode(ps_script.encode("utf-16le")).decode("ascii")
    result = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-EncodedCommand", encoded],
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        raise ResumeSourceError(result.stderr.strip() or result.stdout.strip() or "Microsoft Word photo insertion failed.")


def render_docx(resume: Resume, output: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt

    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(11)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(11.5)
    section.right_margin = Mm(11.5)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)

    use_floating_photo = bool(resume.photo and os.name == "nt")
    use_inline_photo = bool(resume.photo and not use_floating_photo)

    if use_inline_photo:
        header = document.add_table(rows=1, cols=2)
        header.autofit = False
        left_cell, right_cell = header.rows[0].cells
        left_cell.width = Mm(150)
        right_cell.width = Mm(30)
        title = left_cell.paragraphs[0]
    else:
        title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run(resume.name)
    set_run_font(title_run, "Microsoft YaHei", 22, True, resume.accent)
    if resume.target_role:
        target_run = title.add_run(f"  {resume.target_role}")
        set_run_font(target_run, "Microsoft YaHei", 11, True, "606A75")

    if resume.contact:
        contact = left_cell.add_paragraph("  |  ".join(resume.contact)) if use_inline_photo else document.add_paragraph("  |  ".join(resume.contact))
        style_paragraph(contact, size_pt=9.5)
    if use_inline_photo and resume.photo:
        photo = right_cell.paragraphs[0]
        photo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        photo.add_run().add_picture(str(resume.photo), width=Mm(24), height=Mm(28))

    for section_data in resume.sections:
        add_colored_heading(document, section_data.title, 2, resume.accent)
        for paragraph_text in section_data.paragraphs:
            paragraph = document.add_paragraph(paragraph_text)
            style_paragraph(paragraph)
        for item in section_data.items:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Mm(5)
            paragraph.paragraph_format.first_line_indent = Mm(-2.5)
            paragraph.add_run(item)
            style_paragraph(paragraph)
        for entry in section_data.entries:
            if entry.heading:
                add_colored_heading(document, entry.heading, 3, resume.accent)
            for paragraph_text in entry.paragraphs:
                paragraph = document.add_paragraph(paragraph_text)
                style_paragraph(paragraph)
            for bullet in entry.bullets:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.left_indent = Mm(5)
                paragraph.paragraph_format.first_line_indent = Mm(-2.5)
                paragraph.add_run(bullet)
                style_paragraph(paragraph)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    if use_floating_photo and resume.photo:
        add_floating_photo_with_word(output, resume.photo)


def export_pdf_with_word(docx_path: Path, pdf_path: Path) -> None:
    if os.name != "nt":
        raise ResumeSourceError("PDF export requires Microsoft Word on Windows.")
    safe_docx = str(docx_path).replace("'", "''")
    safe_pdf = str(pdf_path).replace("'", "''")
    ps_script = f"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {{
  $doc = $word.Documents.Open('{safe_docx}')
  $doc.SaveAs([ref] '{safe_pdf}', [ref] 17)
  $doc.Close([ref] $false)
}} finally {{
  $word.Quit()
}}
"""
    encoded = base64.b64encode(ps_script.encode("utf-16le")).decode("ascii")
    result = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-EncodedCommand", encoded],
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        raise ResumeSourceError(result.stderr.strip() or result.stdout.strip() or "Microsoft Word PDF export failed.")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate A4 resume HTML, DOCX, and PDF from YAML + Markdown.")
    parser.add_argument("--yaml", type=Path, required=True, help="resume.yaml source")
    parser.add_argument("--markdown", type=Path, help="optional resume.md source")
    parser.add_argument("--outdir", type=Path, default=Path("outputs"), help="output directory")
    parser.add_argument("--stem", help="output filename stem without extension")
    parser.add_argument("--formats", nargs="+", default=["html", "docx", "pdf"], choices=["html", "docx", "pdf"])
    parser.add_argument("--css", type=Path, default=DEFAULT_CSS, help="CSS template for HTML")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    data = load_yaml_like(args.yaml)
    data["_source_dir"] = str(args.yaml.resolve().parent)
    markdown_sections = parse_markdown_sections(args.markdown)
    resume = build_resume(data, markdown_sections)
    stem = args.stem or default_stem(resume)
    args.outdir.mkdir(parents=True, exist_ok=True)

    html_path = args.outdir / f"{stem}.html"
    docx_path = args.outdir / f"{stem}.docx"
    pdf_path = args.outdir / f"{stem}.pdf"

    if "html" in args.formats:
        render_html(resume, html_path, args.css)
        print(f"HTML: {html_path}")
    if "docx" in args.formats or "pdf" in args.formats:
        render_docx(resume, docx_path)
        print(f"DOCX: {docx_path}")
    if "pdf" in args.formats:
        export_pdf_with_word(docx_path, pdf_path)
        print(f"PDF: {pdf_path}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except ResumeSourceError as exc:
        print(f"resume-publisher-cn: {exc}", file=sys.stderr)
        raise SystemExit(2)
