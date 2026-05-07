#!/usr/bin/env python
"""Extract plain text from a DOCX resume for source conversion."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document


def extract_docx_text(path: Path) -> str:
    document = Document(str(path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                lines.append(line)
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract plain text from a DOCX resume.")
    parser.add_argument("docx", type=Path, help="Input .docx file")
    parser.add_argument("--output", "-o", type=Path, help="Optional text output path")
    args = parser.parse_args()

    text = extract_docx_text(args.docx)
    if not text.strip():
        print(
            "No extractable text found. The DOCX may be image-based or use complex drawing objects.",
            file=sys.stderr,
        )
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text, encoding="utf-8")
    else:
        print(text)


if __name__ == "__main__":
    main()
